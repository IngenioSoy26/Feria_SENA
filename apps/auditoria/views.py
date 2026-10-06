import io
from datetime import datetime

from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Q
from django.http import HttpResponse
from django.template.response import TemplateResponse
from django.views import View
from django.utils.dateparse import parse_datetime

from apps.auditoria.models import AuditLog
from apps.core.mixins import RoleRequiredMixin
from apps.usuarios.models import Usuario


class AuditoriaListView(RoleRequiredMixin, LoginRequiredMixin, View):
    roles_requeridos = ['ADMINISTRADOR']
    template_name = 'auditoria/lista.html'
    paginate_by = 25

    def get(self, request, *args, **kwargs):
        usuario_filtro = request.GET.get('usuario', '').strip()
        accion_filtro = request.GET.get('accion', '').strip()
        modulo_filtro = request.GET.get('modulo', '').strip()
        fecha_desde = request.GET.get('fecha_desde', '').strip()
        fecha_hasta = request.GET.get('fecha_hasta', '').strip()
        exportar = request.GET.get('exportar', '').strip().lower()

        qs = AuditLog.objects.select_related('usuario').all()

        if usuario_filtro:
            try:
                uid = int(usuario_filtro)
                qs = qs.filter(Q(usuario_id=uid) | Q(usuario__username__icontains=usuario_filtro))
            except ValueError:
                qs = qs.filter(
                    Q(usuario__username__icontains=usuario_filtro)
                    | Q(usuario__nombres__icontains=usuario_filtro)
                    | Q(usuario__apellidos__icontains=usuario_filtro)
                )
        if accion_filtro:
            qs = qs.filter(accion=accion_filtro)
        if modulo_filtro:
            qs = qs.filter(modulo__icontains=modulo_filtro)
        if fecha_desde:
            try:
                desde = parse_datetime(fecha_desde) if 'T' in fecha_desde or ' ' in fecha_desde else datetime.strptime(fecha_desde, '%Y-%m-%d')
                qs = qs.filter(fecha_hora__gte=desde)
            except ValueError:
                pass
        if fecha_hasta:
            try:
                hasta = parse_datetime(fecha_hasta) if 'T' in fecha_hasta or ' ' in fecha_hasta else datetime.strptime(fecha_hasta + ' 23:59:59', '%Y-%m-%d %H:%M:%S')
                qs = qs.filter(fecha_hora__lte=hasta)
            except ValueError:
                pass

        if exportar == 'excel':
            return self._exportar_excel(qs)

        total = qs.count()
        try:
            pagina = int(request.GET.get('page', 1))
        except ValueError:
            pagina = 1
        if pagina < 1:
            pagina = 1
        offset = (pagina - 1) * self.paginate_by
        limite = offset + self.paginate_by
        registros = list(qs[offset:limite])
        total_paginas = max(1, (total + self.paginate_by - 1) // self.paginate_by)
        if pagina > total_paginas:
            pagina = total_paginas

        usuarios = Usuario.objects.filter(activo=True).order_by('apellidos', 'nombres')
        acciones = [c[0] for c in AuditLog.ACCIONES]

        querystring_sin_page = request.GET.copy()
        querystring_sin_page.pop('page', None)
        querystring_limpia = querystring_sin_page.urlencode()

        context = {
            'registros': registros,
            'pagina_actual': pagina,
            'total_paginas': total_paginas,
            'total': total,
            'tiene_anterior': pagina > 1,
            'tiene_siguiente': pagina < total_paginas,
            'pagina_anterior': pagina - 1,
            'pagina_siguiente': pagina + 1,
            'filtro_usuario': usuario_filtro,
            'filtro_accion': accion_filtro,
            'filtro_modulo': modulo_filtro,
            'filtro_fecha_desde': fecha_desde,
            'filtro_fecha_hasta': fecha_hasta,
            'acciones': acciones,
            'usuarios': usuarios,
            'querystring_limpia': querystring_limpia,
            'rango_paginas': self._rango_paginas(pagina, total_paginas),
        }
        return TemplateResponse(request, self.template_name, context)

    def _rango_paginas(self, actual, total, vecinos=2):
        paginas = set()
        for p in range(max(1, actual - vecinos), min(total, actual + vecinos) + 1):
            paginas.add(p)
        paginas.add(1)
        paginas.add(total)
        return sorted(paginas)

    def _exportar_excel(self, qs):
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = 'Auditoria'

        encabezados = [
            'Fecha y Hora', 'Usuario', 'Rol Usuario',
            'Acción', 'Módulo', 'Entidad', 'ID Entidad',
            'IP', 'User Agent',
        ]
        ws.append(encabezados)

        header_fill = PatternFill(start_color='39A900', end_color='39A900', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for log in qs.iterator():
            usuario_nombre = ''
            rol = ''
            if log.usuario:
                usuario_nombre = log.usuario.nombre_completo if hasattr(log.usuario, 'nombre_completo') else log.usuario.username
                rol = getattr(log.usuario, 'rol_sistema', '')
            row = [
                log.fecha_hora.strftime('%Y-%m-%d %H:%M:%S') if log.fecha_hora else '',
                usuario_nombre,
                rol,
                log.get_accion_display() if log.accion else '',
                log.modulo or '',
                log.entidad or '',
                str(log.id_entidad) if log.id_entidad else '',
                log.ip or '',
                log.user_agent or '',
            ]
            ws.append(row)

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_len:
                        max_len = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'auditoria_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
