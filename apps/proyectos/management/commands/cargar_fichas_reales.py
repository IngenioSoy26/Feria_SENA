import os
import re
from pathlib import Path

from django.core.management.base import BaseCommand
from django.db import transaction

from apps.instituciones.models import InstitucionEducativa
from apps.programas.models import ProgramaTecnico
from apps.eventos.models import Evento, TipoIdentificacion
from apps.personas.models import Persona
from apps.instructores.models import Instructor
from apps.proyectos.models import Proyecto

COL_FICHA = 0
COL_PROGRAMA = 1
COL_INSTITUCION = 2
COL_GRADO = 3
COL_MUNICIPIO = 4
COL_LIDER = 5


def _separar_nombres_apellidos(texto):
    texto = (texto or '').strip()
    if not texto:
        return 'Sin Nombre', 'Sin Apellido'
    partes = re.split(r'\s+', texto)
    if len(partes) == 1:
        return partes[0], 'Sin Apellido'
    if len(partes) == 2:
        return partes[0], partes[1]
    if len(partes) == 3:
        return partes[0], ' '.join(partes[1:])
    return ' '.join(partes[:2]), ' '.join(partes[2:])


class Command(BaseCommand):
    help = 'Carga las 49 fichas reales desde INFORMACIÓN_POR_FICHA.xlsx al catálogo (Instituciones, Programas, Instructores, Proyectos).'

    def add_arguments(self, parser):
        parser.add_argument(
            '--excel',
            type=str,
            default=None,
            help='Ruta al archivo Excel (default: raiz del proyecto INFORMACIÓN_POR_FICHA.xlsx)',
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            self.stderr.write(self.style.ERROR(f'openpyxl no instalado: {exc}'))
            return

        excel_path = options.get('excel')
        if not excel_path:
            raiz = Path(os.path.dirname(os.path.dirname(os.path.dirname(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            ))))
            excel_path = raiz / 'INFORMACIÓN_POR_FICHA.xlsx'
        else:
            excel_path = Path(excel_path)

        if not excel_path.exists():
            self.stderr.write(self.style.ERROR(f'No existe el Excel: {excel_path}'))
            return

        self.stdout.write('=' * 72)
        self.stdout.write('  CARGA DE FICHAS REALES - SISTEMA FERIA PROYECTOS SENA 2026')
        self.stdout.write('=' * 72)
        self.stdout.write(f'Archivo Excel: {excel_path}')

        wb = load_workbook(filename=str(excel_path), data_only=True)
        ws = wb.active
        if 'Datos' in wb.sheetnames:
            ws = wb['Datos']

        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            self.stderr.write(self.style.ERROR('Excel vacío.'))
            return

        header = filas[0]
        self.stdout.write(f'Columnas detectadas: {header}')

        tipo_cc, _ = TipoIdentificacion.objects.get_or_create(
            codigo='CC',
            defaults={'nombre': 'Cédula de Ciudadanía', 'activo': True},
        )

        evento, creado_evento = Evento.objects.get_or_create(
            nombre='FERIA NACIONAL DE PROYECTOS PRODUCTIVOS SENA 2026',
            defaults={
                'fecha_inicio': '2026-11-20',
                'fecha_fin': '2026-11-22',
                'lugar': 'Coliseo Cubierto de Riohacha',
                'municipio': 'Riohacha',
                'regional': 'Guajira',
                'estado': 'ACTIVO',
                'activo': True,
            },
        )
        if creado_evento:
            self.stdout.write(self.style.SUCCESS(f'[+] Evento creado: {evento.nombre}'))
        else:
            self.stdout.write(f'[=] Evento existente: {evento.nombre}')

        c_instituciones_nuevas = 0
        c_programas_nuevos = 0
        c_instructores_nuevos = 0
        c_proyectos_nuevos = 0
        c_proyectos_actualizados = 0
        cedula_seq = 90000000

        with transaction.atomic():
            for idx, fila in enumerate(filas[1:], start=1):
                if not fila or all(v is None or (isinstance(v, str) and not v.strip()) for v in fila):
                    continue

                ficha = fila[COL_FICHA]
                nombre_programa = fila[COL_PROGRAMA]
                nombre_institucion = fila[COL_INSTITUCION]
                grado = fila[COL_GRADO]
                municipio = fila[COL_MUNICIPIO]
                lider = fila[COL_LIDER]

                if ficha is None or nombre_programa is None or nombre_institucion is None:
                    self.stdout.write(self.style.WARNING(f'[?] Fila {idx}: sin ficha/prog/institucion, salta.'))
                    continue

                codigo_ficha = str(ficha).strip()
                nombre_programa = str(nombre_programa).strip()
                nombre_institucion = str(nombre_institucion).strip()
                municipio = str(municipio or 'Riohacha').strip() or 'Riohacha'
                grado = str(grado or '11').strip() or '11'
                lider = str(lider or '').strip()

                codigo_ie = f"INS-{codigo_ficha}"
                institucion, creada_ie = InstitucionEducativa.objects.get_or_create(
                    nombre=nombre_institucion,
                    defaults={
                        'codigo': codigo_ie,
                        'municipio': municipio,
                        'secretaria_educacion': f'Secretaría Educación {municipio}',
                    },
                )
                if creada_ie:
                    c_instituciones_nuevas += 1

                codigo_prog = f"PROG-{codigo_ficha}"
                programa, creado_prog = ProgramaTecnico.objects.get_or_create(
                    nombre=nombre_programa,
                    defaults={'codigo': codigo_prog},
                )
                if creado_prog:
                    c_programas_nuevos += 1

                instructor_obj = None
                if lider:
                    nombres_lider, apellidos_lider = _separar_nombres_apellidos(lider)
                    cedula_seq += 1
                    cedula_lider = f'{cedula_seq}'
                    persona_lider, creada_persona = Persona.objects.get_or_create(
                        tipo_identificacion=tipo_cc,
                        numero_identificacion=cedula_lider,
                        defaults={
                            'nombres': nombres_lider,
                            'apellidos': apellidos_lider,
                            'correo': f'lider_{codigo_ficha.lower()}@senaferia.edu.co',
                            'telefono': '',
                            'tipo_persona': 'INSTRUCTOR',
                        },
                    )
                    instructor_obj, creado_instr = Instructor.objects.get_or_create(
                        persona=persona_lider,
                    )
                    if not instructor_obj.programas.filter(pk=programa.pk).exists():
                        instructor_obj.programas.add(programa)
                    if creado_instr or creada_persona:
                        c_instructores_nuevos += 1

                nombre_proyecto = f'Proyecto {nombre_programa} - {codigo_ficha}'
                proyecto, creado_proy = Proyecto.objects.update_or_create(
                    evento=evento,
                    codigo=codigo_ficha,
                    defaults={
                        'nombre': nombre_proyecto[:250],
                        'descripcion': (
                            f'Ficha {codigo_ficha}. Programa: {nombre_programa}. '
                            f'Institución Educativa: {nombre_institucion}. '
                            f'Grado: {grado}. Municipio: {municipio}. '
                            f'Lider/instructor: {lider or "Sin asignar"}.'
                        ),
                        'institucion': institucion,
                        'programa': programa,
                        'instructor_responsable': instructor_obj,
                        'estado': 'APROBADO',
                    },
                )
                if creado_proy:
                    c_proyectos_nuevos += 1
                else:
                    c_proyectos_actualizados += 1

        self.stdout.write('')
        self.stdout.write(self.style.SUCCESS('✔ CARGA TERMINADA'))
        self.stdout.write(f'  Instituciones nuevas:     {c_instituciones_nuevas}')
        self.stdout.write(f'  Programas nuevos:         {c_programas_nuevos}')
        self.stdout.write(f'  Instructores + Personas:  {c_instructores_nuevos}')
        self.stdout.write(f'  Proyectos nuevos:         {c_proyectos_nuevos}')
        self.stdout.write(f'  Proyectos actualizados:   {c_proyectos_actualizados}')
        self.stdout.write(f'  Evento (ACTIVO): {evento.nombre}')
        self.stdout.write('')
        self.stdout.write(
            f'Total en BD: Instituciones={InstitucionEducativa.objects.count()} | '
            f'Programas={ProgramaTecnico.objects.count()} | '
            f'Instructores={Instructor.objects.count()} | '
            f'Proyectos={Proyecto.objects.filter(evento=evento).count()} (de {Proyecto.objects.count()} totales)'
        )
        self.stdout.write('=' * 72)
